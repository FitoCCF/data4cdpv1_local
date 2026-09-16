# -*- coding: utf-8 -*-
"""
scripts/courier_excel_clone/provider.py
=======================================
Extracción de datos para la réplica exacta de 'Courier_AUTO-C2.xlsm':
- Modo DEV: Lee directamente de 'Courier_AUTO-C2.xlsm' (hojas Portada, TABLA, TABLA_M).
- Modo PROD: Conecta a PiGateway (127.0.0.1:5000) y calcula las métricas estándar de Excel.
"""

from datetime import datetime, date
import os
from typing import Dict, Any, List, Optional
import numpy as np
import openpyxl
import pandas as pd

from scripts.courier_excel_clone.config import (
    STREAMS_COBRE,
    STREAMS_MOLY,
    INTERMEDIATE_STREAMS_COBRE,
    INTERMEDIATE_STREAMS_MOLY,
)


class ExcelCloneProvider:
    """
    Proveedor de datos que extrae los valores directamente de las hojas
    'Portada', 'TABLA' y 'TABLA_M' del archivo Excel original.
    """

    def __init__(self, excel_path: str = "data/raw/Courier_AUTO-C2.xlsm"):
        # Si la ruta es relativa y estamos en scripts/, buscamos en ../data/raw/...
        if not os.path.exists(excel_path):
            alt_path = os.path.join(os.path.dirname(__file__), "../../data/raw/Courier_AUTO-C2.xlsm")
            if os.path.exists(alt_path):
                excel_path = os.path.abspath(alt_path)

        if not os.path.exists(excel_path):
            raise FileNotFoundError(f"No se encontró el archivo Excel en: {excel_path}")

        self.excel_path = excel_path
        print(f"[ExcelCloneProvider] Cargando libro: {self.excel_path}")
        self.wb = openpyxl.load_workbook(self.excel_path, data_only=True)
        print("[ExcelCloneProvider] Libro cargado en memoria exitosamente.")

    def get_portada_data(self) -> Dict[str, Any]:
        """
        Extrae exactamente todas las celdas de la hoja Portada tal cual las muestra Excel.
        """
        ws = self.wb["Portada"]

        def _val(coord: str, default: str = "Sin dato"):
            v = ws[coord].value
            if v is None or v == "":
                return default
            if isinstance(v, float):
                if np.isnan(v):
                    return default
                return v
            return v

        def _fmt(val, decimals: int = 3, is_pct: bool = False):
            if val == "Sin dato" or val is None:
                return "Sin dato"
            try:
                fval = float(val)
                if np.isnan(fval):
                    return "Sin dato"
                txt = f"{fval:.{decimals}f}"
                if is_pct:
                    txt += "%"
                return txt
            except (ValueError, TypeError):
                return str(val)

        # Fecha de evaluación en P5
        raw_date = ws["P5"].value
        if isinstance(raw_date, (datetime, date)):
            date_str = raw_date.strftime("%d-%b-%Y")
        else:
            date_str = str(raw_date)

        # ----------------------------------------------------------------------
        # 1. PLANTA DE COBRE C2
        # ----------------------------------------------------------------------
        cobre_data = {
            "title": "Comparación Courier Planta de Cobre C2 - Laboratorio",
            "date_str": date_str,
            "streams": [
                {
                    "name": "Alimentación Rougher Promedio",
                    "elements": [
                        {
                            "name": "%Cu",
                            "turn_a": {"cour": _val("E11"), "lab": _val("F11")},
                            "turn_b": {"cour": _val("G11"), "lab": _val("H11")},
                            "dif_a": _val("E12"), "pct_a": _val("F12"),
                            "dif_b": _val("G12"), "pct_b": _val("H12"),
                            "dec": 3,
                        },
                        {
                            "name": "%Mo",
                            "turn_a": {"cour": _val("I11"), "lab": _val("J11")},
                            "turn_b": {"cour": _val("K11"), "lab": _val("L11")},
                            "dif_a": _val("I12"), "pct_a": _val("J12"),
                            "dif_b": _val("K12"), "pct_b": _val("L12"),
                            "dec": 3,
                        },
                        {
                            "name": "%Fe",
                            "turn_a": {"cour": _val("M11"), "lab": _val("N11")},
                            "turn_b": {"cour": _val("O11"), "lab": _val("P11")},
                            "dif_a": _val("M12"), "pct_a": _val("N12"),
                            "dif_b": _val("O12"), "pct_b": _val("P12"),
                            "dec": 3,
                        },
                        {
                            "name": "%Zn",
                            "turn_a": {"cour": _val("E17"), "lab": _val("F17")},
                            "turn_b": {"cour": _val("G17"), "lab": _val("H17")},
                            "dif_a": _val("E18"), "pct_a": _val("F18"),
                            "dif_b": _val("G18"), "pct_b": _val("H18"),
                            "dec": 3,
                        },
                        {
                            "name": "%Ox",
                            "turn_a": {"cour": _val("I17"), "lab": _val("J17")},
                            "turn_b": {"cour": _val("K17"), "lab": _val("L17")},
                            "dif_a": _val("I18"), "pct_a": _val("J18"),
                            "dif_b": _val("K18"), "pct_b": _val("L18"),
                            "dec": 3,
                        },
                    ]
                },
                {
                    "name": "Cola Final",
                    "elements": [
                        {
                            "name": "%Cu",
                            "turn_a": {"cour": _val("E24"), "lab": _val("F24")},
                            "turn_b": {"cour": _val("G24"), "lab": _val("H24")},
                            "dif_a": _val("E25"), "pct_a": _val("F25"),
                            "dif_b": _val("G25"), "pct_b": _val("H25"),
                            "dec": 3,
                        },
                        {
                            "name": "%Mo",
                            "turn_a": {"cour": _val("I24"), "lab": _val("J24")},
                            "turn_b": {"cour": _val("K24"), "lab": _val("L24")},
                            "dif_a": _val("I25"), "pct_a": _val("J25"),
                            "dif_b": _val("K25"), "pct_b": _val("L25"),
                            "dec": 3,
                        },
                        {
                            "name": "%Fe",
                            "turn_a": {"cour": _val("M24"), "lab": _val("N24")},
                            "turn_b": {"cour": _val("O24"), "lab": _val("P24")},
                            "dif_a": _val("M25"), "pct_a": _val("N25"),
                            "dif_b": _val("O25"), "pct_b": _val("P25"),
                            "dec": 3,
                        },
                    ]
                },
                {
                    "name": "Concentrado Final",
                    "elements": [
                        {
                            "name": "%Cu",
                            "turn_a": {"cour": _val("E31"), "lab": _val("F31")},
                            "turn_b": {"cour": _val("G31"), "lab": _val("H31")},
                            "dif_a": _val("E32"), "pct_a": _val("F32"),
                            "dif_b": _val("G32"), "pct_b": _val("H32"),
                            "dec": 2,
                        },
                        {
                            "name": "%Mo",
                            "turn_a": {"cour": _val("I31"), "lab": _val("J31")},
                            "turn_b": {"cour": _val("K31"), "lab": _val("L31")},
                            "dif_a": _val("I32"), "pct_a": _val("J32"),
                            "dif_b": _val("K32"), "pct_b": _val("L32"),
                            "dec": 3,
                        },
                        {
                            "name": "%Fe",
                            "turn_a": {"cour": _val("M31"), "lab": _val("N31")},
                            "turn_b": {"cour": _val("O31"), "lab": _val("P31")},
                            "dif_a": _val("M32"), "pct_a": _val("N32"),
                            "dif_b": _val("O32"), "pct_b": _val("P32"),
                            "dec": 2,
                        },
                        {
                            "name": "%Ins",
                            "turn_a": {"cour": _val("E37"), "lab": _val("F37")},
                            "turn_b": {"cour": _val("G37"), "lab": _val("H37")},
                            "dif_a": _val("E38"), "pct_a": _val("F38"),
                            "dif_b": _val("G38"), "pct_b": _val("H38"),
                            "dec": 2,
                        },
                        {
                            "name": "%Ox",
                            "turn_a": {"cour": _val("I37"), "lab": _val("J37")},
                            "turn_b": {"cour": _val("K37"), "lab": _val("L37")},
                            "dif_a": _val("I38"), "pct_a": _val("J38"),
                            "dif_b": _val("K38"), "pct_b": _val("L38"),
                            "dec": 3,
                        },
                    ]
                },
            ],
            "recovery": [
                {
                    "name": "Recuperación Cu",
                    "turn_a": {"cour": _val("E44"), "lab": _val("F44")},
                    "turn_b": {"cour": _val("G44"), "lab": _val("H44")},
                },
                {
                    "name": "Recuperación Mo",
                    "turn_a": {"cour": _val("I44"), "lab": _val("J44")},
                    "turn_b": {"cour": _val("K44"), "lab": _val("L44")},
                },
            ],
            "stdev_table": {
                "title": "DIFERENCIA LABORATORIO vs COURIER - DESVIACION ESTANDAR (%)",
                "label": "ULTIMOS 31 DIAS",
                "columns": [
                    ("ALIM. Cu", _val("G49")),
                    ("ALIM. Mo", _val("H49")),
                    ("ALIM. Fe", _val("I49")),
                    ("COLA Cu", _val("J49")),
                    ("COLA Mo", _val("K49")),
                    ("COLA Fe", _val("L49")),
                    ("CONC. Cu", _val("M49")),
                    ("CONC. Mo", _val("N49")),
                    ("CONC. Fe", _val("O49")),
                ]
            },
            "intermediate_streams": []
        }

        # Flujos intermedios Cobre (Filas 53 a 72)
        for r in range(53, 73, 2):
            stream_name = ws.cell(r, 4).value
            # Guardia A
            hrs_a = ws.cell(r, 6).value
            op_a = ws.cell(r, 7).value
            ensayo1_a = ws.cell(r, 8).value
            min1_a = ws.cell(r, 9).value
            max1_a = ws.cell(r, 10).value
            avg1_a = ws.cell(r, 11).value

            ensayo2_a = ws.cell(r + 1, 8).value
            min2_a = ws.cell(r + 1, 9).value
            max2_a = ws.cell(r + 1, 10).value
            avg2_a = ws.cell(r + 1, 11).value

            # Guardia B
            hrs_b = ws.cell(r, 12).value
            op_b = ws.cell(r, 13).value
            ensayo1_b = ws.cell(r, 14).value
            min1_b = ws.cell(r, 15).value
            max1_b = ws.cell(r, 16).value
            avg1_b = ws.cell(r, 17).value

            ensayo2_b = ws.cell(r + 1, 14).value
            min2_b = ws.cell(r + 1, 15).value
            max2_b = ws.cell(r + 1, 16).value
            avg2_b = ws.cell(r + 1, 17).value

            cobre_data["intermediate_streams"].append({
                "stream": stream_name,
                "guardia_a": {
                    "hrs": hrs_a, "op": op_a,
                    "e1": {"name": ensayo1_a, "min": min1_a, "max": max1_a, "avg": avg1_a},
                    "e2": {"name": ensayo2_a, "min": min2_a, "max": max2_a, "avg": avg2_a},
                },
                "guardia_b": {
                    "hrs": hrs_b, "op": op_b,
                    "e1": {"name": ensayo1_b, "min": min1_b, "max": max1_b, "avg": avg1_b},
                    "e2": {"name": ensayo2_b, "min": min2_b, "max": max2_b, "avg": avg2_b},
                }
            })

        # ----------------------------------------------------------------------
        # 2. PLANTA DE MOLIBDENO C2
        # ----------------------------------------------------------------------
        moly_data = {
            "title": "Comparación Courier Planta de Moly C2 - Laboratorio",
            "streams": [
                {
                    "name": "Alimentación Rougher",
                    "elements": [
                        {
                            "name": "%Cu",
                            "turn_a": {"cour": _val("E84"), "lab": _val("F84")},
                            "turn_b": {"cour": _val("G84"), "lab": _val("H84")},
                            "dif_a": _val("E85"), "pct_a": _val("F85"),
                            "dif_b": _val("G85"), "pct_b": _val("H85"),
                            "dec": 3,
                        },
                        {
                            "name": "%Mo",
                            "turn_a": {"cour": _val("I84"), "lab": _val("J84")},
                            "turn_b": {"cour": _val("K84"), "lab": _val("L84")},
                            "dif_a": _val("I85"), "pct_a": _val("J85"),
                            "dif_b": _val("K85"), "pct_b": _val("L85"),
                            "dec": 3,
                        },
                        {
                            "name": "%Fe",
                            "turn_a": {"cour": _val("M84"), "lab": _val("N84")},
                            "turn_b": {"cour": _val("O84"), "lab": _val("P84")},
                            "dif_a": _val("M85"), "pct_a": _val("N85"),
                            "dif_b": _val("O85"), "pct_b": _val("P85"),
                            "dec": 3,
                        },
                    ]
                },
                {
                    "name": "Cola Rougher Promedio",
                    "elements": [
                        {
                            "name": "%Cu",
                            "turn_a": {"cour": _val("E91"), "lab": _val("F91")},
                            "turn_b": {"cour": _val("G91"), "lab": _val("H91")},
                            "dif_a": _val("E92"), "pct_a": _val("F92"),
                            "dif_b": _val("G92"), "pct_b": _val("H92"),
                            "dec": 3,
                        },
                        {
                            "name": "%Mo",
                            "turn_a": {"cour": _val("I91"), "lab": _val("J91")},
                            "turn_b": {"cour": _val("K91"), "lab": _val("L91")},
                            "dif_a": _val("I92"), "pct_a": _val("J92"),
                            "dif_b": _val("K92"), "pct_b": _val("L92"),
                            "dec": 3,
                        },
                        {
                            "name": "%Fe",
                            "turn_a": {"cour": _val("M91"), "lab": _val("N91")},
                            "turn_b": {"cour": _val("O91"), "lab": _val("P91")},
                            "dif_a": _val("M92"), "pct_a": _val("N92"),
                            "dif_b": _val("O92"), "pct_b": _val("P92"),
                            "dec": 3,
                        },
                    ]
                },
                {
                    "name": "Concentrado Ultima limpieza",
                    "elements": [
                        {
                            "name": "%Cu",
                            "turn_a": {"cour": _val("E98"), "lab": _val("F98")},
                            "turn_b": {"cour": _val("G98"), "lab": _val("H98")},
                            "dif_a": _val("E99"), "pct_a": _val("F99"),
                            "dif_b": _val("G99"), "pct_b": _val("H99"),
                            "dec": 3,
                        },
                        {
                            "name": "%Mo",
                            "turn_a": {"cour": _val("I98"), "lab": _val("J98")},
                            "turn_b": {"cour": _val("K98"), "lab": _val("L98")},
                            "dif_a": _val("I99"), "pct_a": _val("J99"),
                            "dif_b": _val("K99"), "pct_b": _val("L99"),
                            "dec": 2,
                        },
                        {
                            "name": "%Fe",
                            "turn_a": {"cour": _val("M98"), "lab": _val("N98")},
                            "turn_b": {"cour": _val("O98"), "lab": _val("P98")},
                            "dif_a": _val("M99"), "pct_a": _val("N99"),
                            "dif_b": _val("O99"), "pct_b": _val("P99"),
                            "dec": 3,
                        },
                        {
                            "name": "%Ins",
                            "turn_a": {"cour": _val("E104"), "lab": _val("F104")},
                            "turn_b": {"cour": _val("G104"), "lab": _val("H104")},
                            "dif_a": _val("E105"), "pct_a": _val("F105"),
                            "dif_b": _val("G105"), "pct_b": _val("H105"),
                            "dec": 2,
                        },
                        {
                            "name": "%Ox",
                            "turn_a": {"cour": _val("I104"), "lab": _val("J104")},
                            "turn_b": {"cour": _val("K104"), "lab": _val("L104")},
                            "dif_a": _val("I105"), "pct_a": _val("J105"),
                            "dif_b": _val("K105"), "pct_b": _val("L105"),
                            "dec": 3,
                        },
                    ]
                },
            ],
            "recovery": [
                {
                    "name": "Recuperación Mo",
                    "turn_a": {"cour": _val("E110"), "lab": _val("F110")},
                    "turn_b": {"cour": _val("G110"), "lab": _val("H110")},
                }
            ],
            "stdev_table": {
                "title": "DIFERENCIA LABORATORIO vs COURIER - DESVIACION ESTANDAR (%)",
                "label": "ULTIMOS 31 DIAS",
                "columns": [
                    ("ALIM. R. Cu", _val("L110")),
                    ("ALIM. R. Mo", _val("M110")),
                    ("COLA R. Cu", _val("N110")),
                    ("COLA R. Mo", _val("O110")),
                    ("CONC. Cu", _val("P110")),
                    ("CONC. Mo", _val("Q110")),
                ]
            },
            "intermediate_streams": [
                {
                    "stream": "ALIM. ROUGHER",
                    "guardia_a": {
                        "hrs": ws.cell(115, 6).value, "op": ws.cell(115, 7).value,
                        "e1": {"name": "%Cu", "min": ws.cell(115, 9).value, "max": ws.cell(115, 10).value, "avg": ws.cell(115, 11).value},
                        "e2": {"name": "%Mo", "min": ws.cell(116, 9).value, "max": ws.cell(116, 10).value, "avg": ws.cell(116, 11).value},
                    },
                    "guardia_b": {
                        "hrs": ws.cell(115, 12).value, "op": ws.cell(115, 13).value,
                        "e1": {"name": "%Cu", "min": ws.cell(115, 15).value, "max": ws.cell(115, 16).value, "avg": ws.cell(115, 17).value},
                        "e2": {"name": "%Mo", "min": ws.cell(116, 15).value, "max": ws.cell(116, 16).value, "avg": ws.cell(116, 17).value},
                    }
                },
                {
                    "stream": "CONC. 8VA LIMPIEZA",
                    "guardia_a": {
                        "hrs": ws.cell(117, 6).value, "op": ws.cell(117, 7).value,
                        "e1": {"name": "%Cu", "min": ws.cell(117, 9).value, "max": ws.cell(117, 10).value, "avg": ws.cell(117, 11).value},
                        "e2": {"name": "%Mo", "min": ws.cell(118, 9).value, "max": ws.cell(118, 10).value, "avg": ws.cell(118, 11).value},
                    },
                    "guardia_b": {
                        "hrs": ws.cell(117, 12).value, "op": ws.cell(117, 13).value,
                        "e1": {"name": "%Cu", "min": ws.cell(117, 15).value, "max": ws.cell(117, 16).value, "avg": ws.cell(117, 17).value},
                        "e2": {"name": "%Mo", "min": ws.cell(118, 15).value, "max": ws.cell(118, 16).value, "avg": ws.cell(118, 17).value},
                    }
                }
            ]
        }

        return {"cobre": cobre_data, "moly": moly_data}

    def get_series_data(self) -> Dict[str, Any]:
        """
        Extrae las series temporales de 62 turnos de las hojas TABLA y TABLA_M
        junto con los límites y las frecuencias de conformidad.
        """
        ws_t = self.wb["TABLA"]
        ws_tm = self.wb["TABLA_M"]

        shifts_cobre = []
        shifts_moly = []

        # 62 filas en TABLA (filas 4 a 65)
        for r in range(4, 66):
            d_val = ws_t.cell(r, 2).value
            if isinstance(d_val, (datetime, date)):
                d_str = d_val.strftime("%d/%m")
            else:
                d_str = str(d_val) if d_val is not None else f"T{r-3}"
            shifts_cobre.append(d_str)

        for r in range(4, 66):
            d_val = ws_tm.cell(r, 2).value
            if isinstance(d_val, (datetime, date)):
                d_str = d_val.strftime("%d/%m")
            else:
                d_str = str(d_val) if d_val is not None else f"T{r-3}"
            shifts_moly.append(d_str)

        def _get_col_series(ws, col_cour: int, col_lab: int, col_dif: int, col_err: int, tol: float):
            c_vals = []
            l_vals = []
            d_vals = []
            e_vals = []
            sup_vals = []
            inf_vals = []

            for r in range(4, 66):
                cv = ws.cell(r, col_cour).value
                lv = ws.cell(r, col_lab).value
                dv = ws.cell(r, col_dif).value
                ev = ws.cell(r, col_err).value

                def _to_f(v):
                    try:
                        return float(v) if v is not None and v != "" and v != " " else np.nan
                    except:
                        return np.nan

                c_f = _to_f(cv)
                l_f = _to_f(lv)
                d_f = _to_f(dv)
                e_f = _to_f(ev)

                c_vals.append(c_f)
                l_vals.append(l_f)
                d_vals.append(d_f)
                e_vals.append(e_f)

                if not np.isnan(l_f):
                    sup_vals.append(l_f * (1.0 + tol))
                    inf_vals.append(l_f * (1.0 - tol))
                else:
                    sup_vals.append(np.nan)
                    inf_vals.append(np.nan)

            # Clasificación de conformidad
            valid_err = [e for e in e_vals if not np.isnan(e)]
            total = len(valid_err)
            if total > 0:
                n_acep = sum(1 for e in valid_err if e <= (tol * 100))
                n_fuera = total - n_acep
                pct_acep = (n_acep / total) * 100.0
                pct_fuera = (n_fuera / total) * 100.0
            else:
                n_acep, n_fuera, pct_acep, pct_fuera = 0, 0, 0.0, 0.0

            return {
                "cour": np.array(c_vals),
                "lab": np.array(l_vals),
                "dif": np.array(d_vals),
                "err": np.array(e_vals),
                "sup": np.array(sup_vals),
                "inf": np.array(inf_vals),
                "tol": tol,
                "n_acep": n_acep,
                "n_fuera": n_fuera,
                "pct_acep": pct_acep,
                "pct_fuera": pct_fuera,
                "total": total,
            }

        # Extracción Cobre
        cobre_series = {
            "shifts": shifts_cobre,
            "alim_cu": _get_col_series(ws_t, 4, 5, 6, 8, 0.11),
            "alim_mo": _get_col_series(ws_t, 11, 12, 13, 15, 0.11),
            "alim_fe": _get_col_series(ws_t, 18, 19, 20, 21, 0.11),
            "alim_zn": _get_col_series(ws_t, 22, 23, 24, 25, 0.11),
            "cola_cu": _get_col_series(ws_t, 30, 31, 32, 34, 0.20),
            "cola_mo": _get_col_series(ws_t, 37, 38, 39, 41, 0.20),
            "cola_fe": _get_col_series(ws_t, 44, 45, 46, 47, 0.20),
            "conc_cu": _get_col_series(ws_t, 49, 50, 51, 53, 0.06),
            "conc_mo": _get_col_series(ws_t, 56, 57, 58, 60, 0.06),
            "conc_fe": _get_col_series(ws_t, 63, 64, 65, 66, 0.06),
            "conc_ins": _get_col_series(ws_t, 68, 69, 70, 71, 0.10),
        }

        # Extracción Moly
        moly_series = {
            "shifts": shifts_moly,
            "alim_cu": _get_col_series(ws_tm, 4, 5, 6, 8, 0.11),
            "alim_mo": _get_col_series(ws_tm, 11, 12, 13, 15, 0.11),
            "alim_fe": _get_col_series(ws_tm, 18, 19, 20, 22, 0.11),
            "cola_cu": _get_col_series(ws_tm, 25, 26, 27, 29, 0.05),
            "cola_mo": _get_col_series(ws_tm, 32, 33, 34, 36, 0.10),
            "cola_fe": _get_col_series(ws_tm, 39, 40, 41, 43, 0.10),
            "conc_cu": _get_col_series(ws_tm, 46, 47, 48, 50, 0.05),
            "conc_mo": _get_col_series(ws_tm, 53, 54, 55, 57, 0.025),
            "conc_fe": _get_col_series(ws_tm, 60, 61, 62, 64, 0.05),
            "conc_ins": _get_col_series(ws_tm, 67, 68, 69, 71, 0.05),
        }

        return {"cobre": cobre_series, "moly": moly_series}


class PiGatewayCloneProvider:
    """
    Proveedor de datos en línea para la réplica exacta de Excel en modo PROD (WSL2 Debian).
    Se conecta a PiGateway (127.0.0.1:5000), descarga los eventos de los últimos 31 días
    para las corrientes de Cobre y Moly, calcula las diferencias y errores relativos
    de cada turno, y genera las estructuras para la hoja Portada y las gráficas.
    """

    def __init__(self, host: Optional[str] = None, puerto: int = 5000):
        from scripts.courier_demo.provider import PiGatewayProvider
        from scripts.courier_demo.config import STREAMS_COBRE as DEMO_COBRE, STREAMS_MOLY as DEMO_MOLY

        self.demo_cobre = DEMO_COBRE
        self.demo_moly = DEMO_MOLY
        self.pi_prov = PiGatewayProvider(host=host, puerto=puerto)
        self.excel_fallback = ExcelCloneProvider()
        self._cached_series: Optional[Dict[str, Any]] = None
        self._cached_dfs: Optional[Dict[str, Dict[str, pd.DataFrame]]] = None

    def _load_all_dfs(self):
        if self._cached_dfs is not None:
            return self._cached_dfs

        print("[PiGatewayCloneProvider] Consultando datos archivados de 31 días desde PI...")
        cobre_dfs = {}
        for s_id, s_cfg in self.demo_cobre.items():
            cobre_dfs[s_id] = {}
            for elem in s_cfg.elements.keys():
                cobre_dfs[s_id][elem] = self.pi_prov.get_stream_data(s_cfg, elem)

        moly_dfs = {}
        for s_id, s_cfg in self.demo_moly.items():
            moly_dfs[s_id] = {}
            for elem in s_cfg.elements.keys():
                moly_dfs[s_id][elem] = self.pi_prov.get_stream_data(s_cfg, elem)

        self._cached_dfs = {"cobre": cobre_dfs, "moly": moly_dfs}
        return self._cached_dfs

    def get_series_data(self) -> Dict[str, Any]:
        if self._cached_series is not None:
            return self._cached_series

        dfs = self._load_all_dfs()
        cobre_dfs = dfs["cobre"]
        moly_dfs = dfs["moly"]

        # Obtenemos las etiquetas de turnos
        ref_df = cobre_dfs["alim_rougher"]["%Cu"]
        shifts_cobre = [f"{row.date.strftime('%d/%m')} {row.shift}" for row in ref_df.itertuples()]

        ref_m_df = moly_dfs["alim_rougher_moly"]["%Cu"]
        shifts_moly = [f"{row.date.strftime('%d/%m')} {row.shift}" for row in ref_m_df.itertuples()]

        def _build_col_series(df: pd.DataFrame, tol: float) -> Dict[str, Any]:
            c_vals = []
            l_vals = []
            d_vals = []
            e_vals = []
            sup_vals = []
            inf_vals = []

            for row in df.itertuples():
                c_f = float(row.val_cour) if row.val_cour is not None else np.nan
                l_f = float(row.val_lab) if row.val_lab is not None else np.nan

                c_vals.append(c_f)
                l_vals.append(l_f)

                if not np.isnan(l_f) and not np.isnan(c_f):
                    diff = l_f - c_f
                    d_vals.append(diff)
                    if abs(l_f) > 1e-6:
                        err = abs(diff) / l_f * 100.0
                        e_vals.append(err)
                    else:
                        e_vals.append(np.nan)
                else:
                    d_vals.append(np.nan)
                    e_vals.append(np.nan)

                if not np.isnan(l_f):
                    sup_vals.append(l_f * (1.0 + tol))
                    inf_vals.append(l_f * (1.0 - tol))
                else:
                    sup_vals.append(np.nan)
                    inf_vals.append(np.nan)

            valid_err = [e for e in e_vals if not np.isnan(e)]
            total = len(valid_err)
            if total > 0:
                n_acep = sum(1 for e in valid_err if e <= (tol * 100))
                n_fuera = total - n_acep
                pct_acep = (n_acep / total) * 100.0
                pct_fuera = (n_fuera / total) * 100.0
            else:
                n_acep, n_fuera, pct_acep, pct_fuera = 0, 0, 0.0, 0.0

            return {
                "cour": np.array(c_vals),
                "lab": np.array(l_vals),
                "dif": np.array(d_vals),
                "err": np.array(e_vals),
                "sup": np.array(sup_vals),
                "inf": np.array(inf_vals),
                "tol": tol,
                "n_acep": n_acep,
                "n_fuera": n_fuera,
                "pct_acep": pct_acep,
                "pct_fuera": pct_fuera,
                "total": total,
            }

        cobre_series = {
            "shifts": shifts_cobre,
            "alim_cu": _build_col_series(cobre_dfs["alim_rougher"]["%Cu"], 0.11),
            "alim_mo": _build_col_series(cobre_dfs["alim_rougher"]["%Mo"], 0.11),
            "alim_fe": _build_col_series(cobre_dfs["alim_rougher"]["%Fe"], 0.11),
            "alim_zn": _build_col_series(cobre_dfs["alim_rougher"]["%Zn"], 0.11),
            "cola_cu": _build_col_series(cobre_dfs["cola_final"]["%Cu"], 0.20),
            "cola_mo": _build_col_series(cobre_dfs["cola_final"]["%Mo"], 0.20),
            "cola_fe": _build_col_series(cobre_dfs["cola_final"]["%Fe"], 0.20),
            "conc_cu": _build_col_series(cobre_dfs["conc_final"]["%Cu"], 0.06),
            "conc_mo": _build_col_series(cobre_dfs["conc_final"]["%Mo"], 0.06),
            "conc_fe": _build_col_series(cobre_dfs["conc_final"]["%Fe"], 0.06),
            "conc_ins": _build_col_series(cobre_dfs["conc_final"]["%Ins"], 0.10),
        }

        moly_series = {
            "shifts": shifts_moly,
            "alim_cu": _build_col_series(moly_dfs["alim_rougher_moly"]["%Cu"], 0.11),
            "alim_mo": _build_col_series(moly_dfs["alim_rougher_moly"]["%Mo"], 0.11),
            "alim_fe": _build_col_series(moly_dfs["alim_rougher_moly"]["%Fe"], 0.11),
            "cola_cu": _build_col_series(moly_dfs["cola_rougher_moly"]["%Cu"], 0.05),
            "cola_mo": _build_col_series(moly_dfs["cola_rougher_moly"]["%Mo"], 0.10),
            "cola_fe": _build_col_series(moly_dfs["cola_rougher_moly"]["%Fe"], 0.10),
            "conc_cu": _build_col_series(moly_dfs["conc_final_moly"]["%Cu"], 0.05),
            "conc_mo": _build_col_series(moly_dfs["conc_final_moly"]["%Mo"], 0.025),
            "conc_fe": _build_col_series(moly_dfs["conc_final_moly"]["%Fe"], 0.05),
            "conc_ins": _build_col_series(moly_dfs["conc_final_moly"]["%Ins"], 0.05),
        }

        self._cached_series = {"cobre": cobre_series, "moly": moly_series}
        return self._cached_series

    def get_portada_data(self) -> Dict[str, Any]:
        series = self.get_series_data()
        dfs = self._load_all_dfs()

        # Fecha más reciente evaluada
        ref_df = dfs["cobre"]["alim_rougher"]["%Cu"]
        latest_date = ref_df["date"].iloc[-1]
        date_str = latest_date.strftime("%d-%b-%Y")

        def _extract_turn_vals(df: pd.DataFrame, dec: int = 3):
            df_day = df[df["date"] == latest_date]
            row_a = df_day[df_day["shift"] == "A"]
            row_b = df_day[df_day["shift"] == "B"]

            c_a = row_a["val_cour"].values[0] if not row_a.empty else None
            l_a = row_a["val_lab"].values[0] if not row_a.empty else None
            c_b = row_b["val_cour"].values[0] if not row_b.empty else None
            l_b = row_b["val_lab"].values[0] if not row_b.empty else None

            d_a = (l_a - c_a) if (l_a is not None and c_a is not None) else "Sin dato"
            p_a = (abs(d_a) / l_a * 100.0) if (isinstance(d_a, float) and l_a is not None and l_a != 0) else "Sin dato"

            d_b = (l_b - c_b) if (l_b is not None and c_b is not None) else "Sin dato"
            p_b = (abs(d_b) / l_b * 100.0) if (isinstance(d_b, float) and l_b is not None and l_b != 0) else "Sin dato"

            return {
                "turn_a": {"cour": c_a if c_a is not None else "Sin dato", "lab": l_a if l_a is not None else "Sin dato"},
                "turn_b": {"cour": c_b if c_b is not None else "Sin dato", "lab": l_b if l_b is not None else "Sin dato"},
                "dif_a": d_a, "pct_a": p_a,
                "dif_b": d_b, "pct_b": p_b,
                "dec": dec,
            }

        def _calc_stdev(series_dict: Dict[str, Any], key: str):
            arr = series_dict[key]["err"]
            valid = arr[~np.isnan(arr)]
            if len(valid) > 1:
                return float(np.std(valid, ddof=1))
            return "Sin dato"

        stdev_cobre_cols = [
            ("ALIM. Cu", _calc_stdev(series["cobre"], "alim_cu")),
            ("ALIM. Mo", _calc_stdev(series["cobre"], "alim_mo")),
            ("ALIM. Fe", _calc_stdev(series["cobre"], "alim_fe")),
            ("COLA Cu", _calc_stdev(series["cobre"], "cola_cu")),
            ("COLA Mo", _calc_stdev(series["cobre"], "cola_mo")),
            ("COLA Fe", _calc_stdev(series["cobre"], "cola_fe")),
            ("CONC. Cu", _calc_stdev(series["cobre"], "conc_cu")),
            ("CONC. Mo", _calc_stdev(series["cobre"], "conc_mo")),
            ("CONC. Fe", _calc_stdev(series["cobre"], "conc_fe")),
        ]

        stdev_moly_cols = [
            ("ALIM. R. Cu", _calc_stdev(series["moly"], "alim_cu")),
            ("ALIM. R. Mo", _calc_stdev(series["moly"], "alim_mo")),
            ("COLA R. Cu", _calc_stdev(series["moly"], "cola_cu")),
            ("COLA R. Mo", _calc_stdev(series["moly"], "cola_mo")),
            ("CONC. Cu", _calc_stdev(series["moly"], "conc_cu")),
            ("CONC. Mo", _calc_stdev(series["moly"], "conc_mo")),
        ]

        def _rec_2prod(f, c, t):
            try:
                if f is not None and c is not None and t is not None and f != "Sin dato" and c != "Sin dato" and t != "Sin dato":
                    f, c, t = float(f), float(c), float(t)
                    den = f * (c - t)
                    if abs(den) > 1e-6:
                        r = (c * (f - t)) / den * 100.0
                        return max(0.0, min(100.0, r))
                return "Sin dato"
            except:
                return "Sin dato"

        alim_cu = _extract_turn_vals(dfs["cobre"]["alim_rougher"]["%Cu"])
        conc_cu = _extract_turn_vals(dfs["cobre"]["conc_final"]["%Cu"])
        cola_cu = _extract_turn_vals(dfs["cobre"]["cola_final"]["%Cu"])

        rec_cu_a_cour = _rec_2prod(alim_cu["turn_a"]["cour"], conc_cu["turn_a"]["cour"], cola_cu["turn_a"]["cour"])
        rec_cu_a_lab = _rec_2prod(alim_cu["turn_a"]["lab"], conc_cu["turn_a"]["lab"], cola_cu["turn_a"]["lab"])
        rec_cu_b_cour = _rec_2prod(alim_cu["turn_b"]["cour"], conc_cu["turn_b"]["cour"], cola_cu["turn_b"]["cour"])
        rec_cu_b_lab = _rec_2prod(alim_cu["turn_b"]["lab"], conc_cu["turn_b"]["lab"], cola_cu["turn_b"]["lab"])

        alim_mo = _extract_turn_vals(dfs["cobre"]["alim_rougher"]["%Mo"])
        conc_mo = _extract_turn_vals(dfs["cobre"]["conc_final"]["%Mo"])
        cola_mo = _extract_turn_vals(dfs["cobre"]["cola_final"]["%Mo"])

        rec_mo_a_cour = _rec_2prod(alim_mo["turn_a"]["cour"], conc_mo["turn_a"]["cour"], cola_mo["turn_a"]["cour"])
        rec_mo_a_lab = _rec_2prod(alim_mo["turn_a"]["lab"], conc_mo["turn_a"]["lab"], cola_mo["turn_a"]["lab"])
        rec_mo_b_cour = _rec_2prod(alim_mo["turn_b"]["cour"], conc_mo["turn_b"]["cour"], cola_mo["turn_b"]["cour"])
        rec_mo_b_lab = _rec_2prod(alim_mo["turn_b"]["lab"], conc_mo["turn_b"]["lab"], cola_mo["turn_b"]["lab"])

        m_alim_mo = _extract_turn_vals(dfs["moly"]["alim_rougher_moly"]["%Mo"])
        m_conc_mo = _extract_turn_vals(dfs["moly"]["conc_final_moly"]["%Mo"])
        m_cola_mo = _extract_turn_vals(dfs["moly"]["cola_rougher_moly"]["%Mo"])

        m_rec_a_cour = _rec_2prod(m_alim_mo["turn_a"]["cour"], m_conc_mo["turn_a"]["cour"], m_cola_mo["turn_a"]["cour"])
        m_rec_a_lab = _rec_2prod(m_alim_mo["turn_a"]["lab"], m_conc_mo["turn_a"]["lab"], m_cola_mo["turn_a"]["lab"])
        m_rec_b_cour = _rec_2prod(m_alim_mo["turn_b"]["cour"], m_conc_mo["turn_b"]["cour"], m_cola_mo["turn_b"]["cour"])
        m_rec_b_lab = _rec_2prod(m_alim_mo["turn_b"]["lab"], m_conc_mo["turn_b"]["lab"], m_cola_mo["turn_b"]["lab"])

        fallback_portada = self.excel_fallback.get_portada_data()

        cobre_portada = {
            "title": "Comparación Courier Planta de Cobre C2 - Laboratorio",
            "date_str": date_str,
            "streams": [
                {
                    "name": "Alimentación Rougher Promedio",
                    "elements": [
                        {"name": "%Cu", **_extract_turn_vals(dfs["cobre"]["alim_rougher"]["%Cu"], 3)},
                        {"name": "%Mo", **_extract_turn_vals(dfs["cobre"]["alim_rougher"]["%Mo"], 3)},
                        {"name": "%Fe", **_extract_turn_vals(dfs["cobre"]["alim_rougher"]["%Fe"], 3)},
                        {"name": "%Zn", **_extract_turn_vals(dfs["cobre"]["alim_rougher"]["%Zn"], 3)},
                        {"name": "%Ox", "turn_a": {"cour": "Sin dato", "lab": "Sin dato"}, "turn_b": {"cour": "Sin dato", "lab": "Sin dato"}, "dif_a": "Sin dato", "pct_a": "Sin dato", "dif_b": "Sin dato", "pct_b": "Sin dato", "dec": 3},
                    ]
                },
                {
                    "name": "Cola Final",
                    "elements": [
                        {"name": "%Cu", **_extract_turn_vals(dfs["cobre"]["cola_final"]["%Cu"], 3)},
                        {"name": "%Mo", **_extract_turn_vals(dfs["cobre"]["cola_final"]["%Mo"], 3)},
                        {"name": "%Fe", **_extract_turn_vals(dfs["cobre"]["cola_final"]["%Fe"], 3)},
                    ]
                },
                {
                    "name": "Concentrado Final",
                    "elements": [
                        {"name": "%Cu", **_extract_turn_vals(dfs["cobre"]["conc_final"]["%Cu"], 2)},
                        {"name": "%Mo", **_extract_turn_vals(dfs["cobre"]["conc_final"]["%Mo"], 3)},
                        {"name": "%Fe", **_extract_turn_vals(dfs["cobre"]["conc_final"]["%Fe"], 2)},
                        {"name": "%Ins", **_extract_turn_vals(dfs["cobre"]["conc_final"]["%Ins"], 2)},
                        {"name": "%Ox", "turn_a": {"cour": "Sin dato", "lab": "Sin dato"}, "turn_b": {"cour": "Sin dato", "lab": "Sin dato"}, "dif_a": "Sin dato", "pct_a": "Sin dato", "dif_b": "Sin dato", "pct_b": "Sin dato", "dec": 3},
                    ]
                }
            ],
            "recovery": [
                {"name": "Recuperación Cu", "turn_a": {"cour": rec_cu_a_cour, "lab": rec_cu_a_lab}, "turn_b": {"cour": rec_cu_b_cour, "lab": rec_cu_b_lab}},
                {"name": "Recuperación Mo", "turn_a": {"cour": rec_mo_a_cour, "lab": rec_mo_a_lab}, "turn_b": {"cour": rec_mo_b_cour, "lab": rec_mo_b_lab}},
            ],
            "stdev_table": {
                "title": "DIFERENCIA LABORATORIO vs COURIER - DESVIACION ESTANDAR (%)",
                "label": "ULTIMOS 31 DIAS",
                "columns": stdev_cobre_cols,
            },
            "intermediate_streams": fallback_portada["cobre"]["intermediate_streams"]
        }

        moly_portada = {
            "title": "Comparación Courier Planta de Moly C2 - Laboratorio",
            "streams": [
                {
                    "name": "Alimentación Rougher",
                    "elements": [
                        {"name": "%Cu", **_extract_turn_vals(dfs["moly"]["alim_rougher_moly"]["%Cu"], 3)},
                        {"name": "%Mo", **_extract_turn_vals(dfs["moly"]["alim_rougher_moly"]["%Mo"], 3)},
                        {"name": "%Fe", **_extract_turn_vals(dfs["moly"]["alim_rougher_moly"]["%Fe"], 3)},
                    ]
                },
                {
                    "name": "Cola Rougher Promedio",
                    "elements": [
                        {"name": "%Cu", **_extract_turn_vals(dfs["moly"]["cola_rougher_moly"]["%Cu"], 3)},
                        {"name": "%Mo", **_extract_turn_vals(dfs["moly"]["cola_rougher_moly"]["%Mo"], 3)},
                        {"name": "%Fe", **_extract_turn_vals(dfs["moly"]["cola_rougher_moly"]["%Fe"], 3)},
                    ]
                },
                {
                    "name": "Concentrado Ultima limpieza",
                    "elements": [
                        {"name": "%Cu", **_extract_turn_vals(dfs["moly"]["conc_final_moly"]["%Cu"], 3)},
                        {"name": "%Mo", **_extract_turn_vals(dfs["moly"]["conc_final_moly"]["%Mo"], 2)},
                        {"name": "%Fe", **_extract_turn_vals(dfs["moly"]["conc_final_moly"]["%Fe"], 3)},
                        {"name": "%Ins", **_extract_turn_vals(dfs["moly"]["conc_final_moly"]["%Ins"], 2)},
                        {"name": "%Ox", "turn_a": {"cour": "Sin dato", "lab": "Sin dato"}, "turn_b": {"cour": "Sin dato", "lab": "Sin dato"}, "dif_a": "Sin dato", "pct_a": "Sin dato", "dif_b": "Sin dato", "pct_b": "Sin dato", "dec": 3},
                    ]
                }
            ],
            "recovery": [
                {"name": "Recuperación Mo", "turn_a": {"cour": m_rec_a_cour, "lab": m_rec_a_lab}, "turn_b": {"cour": m_rec_b_cour, "lab": m_rec_b_lab}},
            ],
            "stdev_table": {
                "title": "DIFERENCIA LABORATORIO vs COURIER - DESVIACION ESTANDAR (%)",
                "label": "ULTIMOS 31 DIAS",
                "columns": stdev_moly_cols,
            },
            "intermediate_streams": fallback_portada["moly"]["intermediate_streams"]
        }

        return {"cobre": cobre_portada, "moly": moly_portada}
